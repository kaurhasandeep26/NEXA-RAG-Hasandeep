"""File validation, extraction, and conservative text chunking."""
from __future__ import annotations

import io
from dataclasses import dataclass
from pathlib import Path

import fitz  # PyMuPDF
from docx import Document

from src.utils import SUPPORTED_EXTENSIONS, clean_text


class IngestionError(ValueError):
    """Raised when a knowledge file cannot be safely ingested."""


@dataclass
class ExtractedDocument:
    filename: str
    text: str
    page_count: int | None = None


def extract_document(filename: str, content: bytes) -> ExtractedDocument:
    extension = Path(filename).suffix.lower()
    if extension not in SUPPORTED_EXTENSIONS:
        raise IngestionError("Only PDF, TXT, and DOCX files are supported.")
    if not content:
        raise IngestionError("This file is empty.")
    if len(content) > 30 * 1024 * 1024:
        raise IngestionError("File is larger than the 30 MB safety limit.")

    try:
        if extension == ".pdf":
            pdf = fitz.open(stream=content, filetype="pdf")
            pages = [page.get_text("text") for page in pdf]
            text, page_count = "\n\n".join(pages), len(pdf)
            pdf.close()
        elif extension == ".docx":
            document = Document(io.BytesIO(content))
            text = "\n\n".join(p.text for p in document.paragraphs)
            page_count = None
        else:
            text = content.decode("utf-8", errors="replace")
            page_count = None
    except Exception as exc:
        raise IngestionError(f"Could not read '{filename}'. It may be corrupted or protected.") from exc

    text = clean_text(text)
    if len(text) < 40:
        raise IngestionError("No readable text was found in this file. Scanned PDFs need OCR first.")
    return ExtractedDocument(filename=filename, text=text, page_count=page_count)


def split_text(text: str, chunk_size: int, overlap: int) -> list[str]:
    """Split near sentence/paragraph boundaries without losing overlap."""
    if chunk_size < 100:
        raise ValueError("Chunk size must be at least 100 characters.")
    if overlap < 0 or overlap >= chunk_size:
        raise ValueError("Overlap must be non-negative and smaller than chunk size.")
    text = clean_text(text)
    chunks: list[str] = []
    start = 0
    while start < len(text):
        end = min(start + chunk_size, len(text))
        if end < len(text):
            boundary = max(text.rfind("\n", start + chunk_size // 2, end), text.rfind(". ", start + chunk_size // 2, end))
            if boundary > start:
                end = boundary + 1
        chunk = text[start:end].strip()
        if chunk:
            chunks.append(chunk)
        if end >= len(text):
            break
        start = max(end - overlap, start + 1)
    return chunks
